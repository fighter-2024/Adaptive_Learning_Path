"""
生成讨论会 PPT：引导式开场 + 准备好的方案材料
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 颜色
DARK   = RGBColor(0x1A, 0x1A, 0x2E)
BLUE   = RGBColor(0x16, 0x21, 0x3E)
ACCENT = RGBColor(0x0F, 0x34, 0x60)
GREEN  = RGBColor(0x0A, 0x8F, 0x5C)
RED    = RGBColor(0xE9, 0x45, 0x60)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GRAY   = RGBColor(0x88, 0x88, 0x88)
LGRAY  = RGBColor(0xEC, 0xF0, 0xF1)
BG     = RGBColor(0xF8, 0xF9, 0xFA)

# ── 辅助函数 ──
def add_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_title_bar(slide, text, y=Inches(0.3)):
    """顶部标题栏"""
    from pptx.util import Inches, Pt
    shape = slide.shapes.add_shape(
        1, Inches(0), y, prs.slide_width, Inches(1.1))  # MSO_SHAPE.RECTANGLE = 1
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(32)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.8)
    tf.margin_top = Inches(0.15)

def add_body_text(slide, text, left, top, width, height, size=Pt(18), color=DARK, bold=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = size
    p.font.color.rgb = color
    p.font.bold = bold
    return tf

def add_bullet_list(slide, items, left, top, width, height, size=Pt(18), color=DARK, spacing=Pt(8)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = size
        p.font.color.rgb = color
        p.space_after = spacing
        p.level = 0
    return tf

def add_question_slide(title, question, sub_items=None):
    """引导性问题幻灯片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide, BG)
    add_title_bar(slide, title)
    
    # 大号问题
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.0), Inches(11), Inches(2.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = question
    p.font.size = Pt(36)
    p.font.color.rgb = RED
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    if sub_items:
        add_bullet_list(slide, sub_items, Inches(2.5), Inches(4.8), Inches(8), Inches(2),
                        size=Pt(20), color=GRAY)
    
    return slide

def add_content_slide(title, content_func):
    """内容幻灯片"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG)
    add_title_bar(slide, title)
    content_func(slide)
    return slide

def add_table(slide, headers, rows, left, top, width, height, col_widths=None):
    """添加表格"""
    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table
    
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    
    # 表头
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
    
    # 数据行
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(14)
                p.font.color.rgb = DARK
                p.alignment = PP_ALIGN.CENTER
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 == 0 else LGRAY
    
    return table_shape


# ═══════════════════════════════════════════
# Slide 1: 封面
# ═══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)

add_body_text(slide, "自适应学习路径规划\n与社交激励系统",
              Inches(1), Inches(1.5), Inches(11), Inches(2.5),
              size=Pt(44), color=WHITE, bold=True)

add_body_text(slide, "项目启动讨论会",
              Inches(1), Inches(4.3), Inches(11), Inches(1),
              size=Pt(28), color=RGBColor(0xE9, 0x45, 0x60), bold=True)

add_body_text(slide, "省级大学生创新创业训练计划项目",
              Inches(1), Inches(5.5), Inches(11), Inches(0.8),
              size=Pt(18), color=GRAY)


# ═══════════════════════════════════════════
# Slide 2-4: 引导性讨论
# ═══════════════════════════════════════════
add_question_slide(
    "Part 1 · 问题定位",
    "我们要解决什么问题？",
    ["- 学生面对大量知识点，不知道该按什么顺序学",
     "- 传统学习平台「千人一面」，缺乏个性化路径",
     "- 学习过程枯燥，缺少同伴激励和成就感",
     "- 教师难以精准诊断每个学生的薄弱点"]
)

add_question_slide(
    "Part 1 · 用户画像",
    "我们的用户是谁？他们需要什么？",
    ["- 学员：需要一条「看得见」的个性化学习路径 + 社交激励",
     "- 教师/管理员：需要一个能诊断学情、管理内容的后台",
     "- 讨论：先聚焦哪类用户？哪门课程？"]
)

add_question_slide(
    "Part 1 · 核心场景",
    "典型使用场景是怎样的？",
    ["- 学员登录 → 做几道诊断题 → 系统给出知识掌握画像",
     "- 系统推荐下一步该学什么、为什么",
     "- 学完知识点 → 做练习题 → 系统更新诊断 → 调整路径",
     "- 查看排行榜 / 好友进度 → 获得激励",
     "- 每周收到 AI 生成的学习周报"])

# ═══════════════════════════════════════════
# Slide 5: 过渡页
# ═══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_body_text(slide, "Part 2",
              Inches(1), Inches(2.0), Inches(11), Inches(1),
              size=Pt(24), color=GRAY)
add_body_text(slide, "基于以上思考，\n我们准备了一套技术方案",
              Inches(1), Inches(3.0), Inches(11), Inches(2),
              size=Pt(38), color=WHITE, bold=True)


# ═══════════════════════════════════════════
# Slide 6: 技术栈
# ═══════════════════════════════════════════
def s6(slide):
    add_table(slide,
        ["层面", "选型", "说明"],
        [
            ["后端", "Python FastAPI", "统一 REST API 服务"],
            ["图数据库", "Neo4j", "知识点关系 & 路径查询"],
            ["关系数据库", "SQL Server", "用户 / 答题 / 社交数据"],
            ["学员前端", "UniApp", "一套代码 → H5 + 微信小程序"],
            ["管理后台", "Vue 3 SPA", "独立工程 admin/"],
            ["认知诊断", "DINA 模型", "Q矩阵 + EM + 贝叶斯后验"],
            ["路径规划", "拓扑排序 + 多指标贪心", "保序 + 距离/难度/掌握概率打分"],
            ["AI 能力", "LLM API (DeepSeek / 通义千问)", "诊断解读 / 错题归因 / 周报"],
        ],
        Inches(1.5), Inches(1.8), Inches(10), Inches(5.2),
        col_widths=[Inches(2.0), Inches(3.2), Inches(4.8)]
    )
add_content_slide("技术栈总览", s6)


# ═══════════════════════════════════════════
# Slide 7: 系统架构图
# ═══════════════════════════════════════════
def s7(slide):
    img_path = os.path.join(os.path.dirname(__file__), "项目结构图.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(
            img_path,
            Inches(1.5), Inches(1.6),
            Inches(10.33), Inches(5.5)
        )
    else:
        add_body_text(slide, "(结构图未找到，请运行 generate_structure_diagram.py)",
                      Inches(1), Inches(3), Inches(11), Inches(1), size=Pt(18), color=RED)
add_content_slide("系统架构总览", s7)


# ═══════════════════════════════════════════
# Slide 8: 项目目录
# ═══════════════════════════════════════════
def s8(slide):
    dirs = [
        "docs/          文档（申报书 / API / 数据库 / 会议）",
        "server/        FastAPI 后端",
        "  └─ app/      config / routers / models / services / db",
        "admin/         Vue 3 管理后台（独立工程）",
        "frontend/      UniApp 学员客户端（H5 + 小程序）",
        "algorithm/     算法实验（DINA + 路径规划 / Jupyter）",
        "kg-data/       知识图谱数据 & 构建脚本",
    ]
    add_bullet_list(slide,
        ["项目目录结构：", ""] + dirs,
        Inches(1.5), Inches(1.8), Inches(10), Inches(5.3),
        size=Pt(17), color=DARK, spacing=Pt(10))
add_content_slide("项目目录结构", s8)


# ═══════════════════════════════════════════
# Slide 9: 学员端功能
# ═══════════════════════════════════════════
def s9(slide):
    add_table(slide,
        ["模块", "功能描述"],
        [
            ["知识点学习 & 答题", "按推荐路径学习知识点，完成配套练习"],
            ["学习路径展示", "可视化当前进度 + 下一步推荐 + 目标距离"],
            ["社交激励", "排行榜 / 打卡 / 好友PK / 成就系统"],
            ["AI 诊断解读", "把认知诊断结果翻译成人话，告诉学员哪里薄弱"],
            ["错题归因", "每道错题给出知识点层面的原因分析"],
            ["学习周报", "每周 AI 生成学习总结 + 进步趋势"],
        ],
        Inches(2.0), Inches(1.8), Inches(9), Inches(5.0),
        col_widths=[Inches(2.5), Inches(6.5)]
    )
add_content_slide("学员端功能 (UniApp)", s9)


# ═══════════════════════════════════════════
# Slide 10: 管理后台功能
# ═══════════════════════════════════════════
def s10(slide):
    add_table(slide,
        ["模块", "功能"],
        [
            ["知识图谱管理", "知识点 / 关系增删改查、图谱可视化"],
            ["题库管理", "题目录入编辑、Q 矩阵标注"],
            ["学生数据看板", "学生列表、认知诊断 α 向量、学习进度"],
            ["社交内容管理", "动态审核、公告发布、排行榜配置、举报处理"],
            ["AI 诊断与周报", "触发 AI 生成解读/周报、历史查看"],
            ["系统配置", "DINA 参数调节、LLM 配置、阈值设置"],
        ],
        Inches(2.0), Inches(1.8), Inches(9), Inches(5.0),
        col_widths=[Inches(2.5), Inches(6.5)]
    )
add_content_slide("管理后台功能 (Vue 3)", s10)


# ═══════════════════════════════════════════
# Slide 11: 核心算法路线
# ═══════════════════════════════════════════
def s11(slide):
    steps = [
        "1.  学生完成一组诊断性题目",
        "      ↓",
        "2.  DINA 模型：EM 算法估计题目参数 (slip / guess)",
        "      ↓",
        "3.  贝叶斯后验推断：得到学生知识掌握向量 α",
        "      ↓",
        "4.  加载 Neo4j 知识图谱 + α 向量",
        "      ↓",
        "5.  拓扑排序确保先修关系 + 多指标贪心打分",
        "      指标：距离目标 / 掌握概率 / 难度 / 预估时长",
        "      ↓",
        "6.  输出个性化推荐学习路径 → 前端展示 + AI 解读",
    ]
    add_bullet_list(slide, steps,
                    Inches(2.0), Inches(1.8), Inches(9), Inches(5.3),
                    size=Pt(18), color=DARK, spacing=Pt(6))
add_content_slide("核心算法路线", s11)


# ═══════════════════════════════════════════
# Slide 12: AI 接入点
# ═══════════════════════════════════════════
def s12(slide):
    items = [
        ("诊断结果解读", "把 α 向量翻译成学员能看懂的语言"),
        ("路径推荐解释", "为什么推荐这个知识点？为什么不推荐那个？"),
        ("错题归因分析", "这道题错了 → 哪个知识点没掌握 → 怎么补"),
        ("学习周报生成", "本周学了多少、正确率、进步趋势、下周建议"),
        ("知识点提取辅助 (可选)", "从教材/讲义自动抽取知识点，辅助建图谱"),
    ]
    for i, (title, desc) in enumerate(items):
        y = Inches(2.0) + Inches(1.0) * i
        add_body_text(slide, f"{title}",
                      Inches(2.0), y, Inches(3.5), Inches(0.7),
                      size=Pt(20), color=RED, bold=True)
        add_body_text(slide, desc,
                      Inches(5.8), y, Inches(6), Inches(0.7),
                      size=Pt(18), color=DARK)
add_content_slide("AI 能力接入点", s12)


# ═══════════════════════════════════════════
# Slide 13: 开发规范
# ═══════════════════════════════════════════
def s13(slide):
    rules = [
        "算法先在 Jupyter Notebook 验证 → 成熟后移植到 server/app/services/",
        "API 统一返回格式：{\"code\": 0, \"data\": ..., \"message\": \"\"}",
        "路由按角色前缀：/api/admin/*  和  /api/student/*",
        "后端代码注释用中文，标识符用英文",
        "所有配置走环境变量 .env，不硬编码密码/密钥",
        "两个前端 (admin + UniApp) 共用同一套 server API",
    ]
    add_bullet_list(slide, rules,
                    Inches(1.5), Inches(2.0), Inches(10), Inches(5),
                    size=Pt(20), color=DARK, spacing=Pt(16))
add_content_slide("开发规范 (已定)", s13)


# ═══════════════════════════════════════════
# Slide 14: 待讨论
# ═══════════════════════════════════════════
add_question_slide(
    "Part 3 · 待讨论事项",
    "接下来需要一起决定",
    [
        "1. 社交激励具体功能范围？排行榜 / 好友 / 小组 / 打卡？",
        "2. admin/ UI 框架选什么？Element Plus / Ant Design Vue？",
        "3. 知识图谱先做哪一门课程？",
        "4. 第一阶段的开发优先级怎么排？",
        "5. 分工 & 时间节点？",
        "",
        "（本页为讨论会历史快照：以上事项后续均已定稿 —",
        "Element Plus / 初中数学 / 排行榜·打卡·成就，详见 docs/）",
    ]
)


# ═══════════════════════════════════════════
# Slide 15: 结束页
# ═══════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_body_text(slide, "感谢参与\n开始讨论吧",
              Inches(1), Inches(2.5), Inches(11), Inches(2.5),
              size=Pt(44), color=WHITE, bold=True)

# ── 保存 ──
out = os.path.join(os.path.dirname(__file__), "讨论会PPT.pptx")
prs.save(out)
print(f"OK: {out}")
print(f"共 {len(prs.slides)} 页")
